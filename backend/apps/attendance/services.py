from collections import defaultdict
from datetime import datetime, time, timedelta
from zoneinfo import ZoneInfo

from django.db import transaction
from openpyxl import load_workbook

from apps.attendance.engine import TimeError, evaluate_times
from apps.attendance.models import AttendanceAdjustment, AttendanceRecord
from apps.common.audit import audit
from apps.common.dates import month_bounds
from apps.common.exceptions import BusinessError
from apps.common.money import hours
from apps.employees.services import (
    WORKING_STATUSES,
    can_manage_attendance,
    employed_on,
    resolve_weekly_offs,
    visible_employees,
)
from apps.organization.models import Holiday, OfficeSettings

VALID_STATUSES = {choice for choice, _label in AttendanceRecord.Status.choices}


def office_now(office=None):
    office = office or OfficeSettings.load()
    try:
        tz = ZoneInfo(office.timezone or "UTC")
    except Exception:
        tz = ZoneInfo("UTC")
    return datetime.now(tz)


def as_time(value):
    if isinstance(value, time):
        return value
    if isinstance(value, str):
        hour, minute, *_rest = value.split(":")
        return time(int(hour), int(minute))
    raise BusinessError("Shift times must use HH:MM.")


def shift_template(employee, office):
    shift = employee.shift if employee.shift_id else None
    if shift is not None:
        tz_name = shift.timezone or (employee.location.timezone if employee.location_id and employee.location.timezone else office.timezone)
        return {
            "shift": shift,
            "start": as_time(shift.start_time),
            "end": as_time(shift.end_time),
            "overnight": shift.is_overnight,
            "break": shift.break_minutes,
            "grace": shift.grace_minutes,
            "early": shift.early_grace_minutes,
            "tz": tz_name or office.timezone,
        }
    return {
        "shift": None,
        "start": time(9, 0),
        "end": time(17, 0),
        "overnight": False,
        "break": 60,
        "grace": office.default_grace_minutes,
        "early": office.early_departure_grace_minutes,
        "tz": office.timezone,
    }


def holiday_dates_for(employee, start, end):
    qs = Holiday.objects.filter(is_active=True, date__range=(start, end)).filter(
        models_location(employee)
    )
    return set(qs.values_list("date", flat=True))


def models_location(employee):
    from django.db.models import Q

    return Q(location__isnull=True) | Q(location_id=employee.location_id)


def leave_index(employee_ids, start, end):
    from apps.leaves.models import LeaveRequest

    grouped = defaultdict(list)
    rows = LeaveRequest.objects.filter(
        employee_id__in=employee_ids,
        status=LeaveRequest.Status.APPROVED,
        start_date__lte=end,
        end_date__gte=start,
    ).select_related("leave_type")
    for row in rows:
        grouped[row.employee_id].append(row)
    return grouped


def leave_dicts(rows):
    return [
        {
            "start": row.start_date,
            "end": row.end_date,
            "day_part": row.day_part,
            "is_paid": row.leave_type.is_paid,
            "status": row.status,
        }
        for row in rows
    ]


def covers_leave(rows, day):
    return any(row.start_date <= day <= row.end_date for row in rows)


def assert_month_editable(day):
    from apps.payroll.models import PayrollPeriod

    if PayrollPeriod.objects.filter(year=day.year, month=day.month, status=PayrollPeriod.Status.FINALIZED).exists():
        raise BusinessError(
            "Payroll for this month is finalized. Reopen it with a reason before changing attendance or leave."
        )
    PayrollPeriod.objects.filter(
        year=day.year,
        month=day.month,
        status__in=[PayrollPeriod.Status.CALCULATED, PayrollPeriod.Status.APPROVED, PayrollPeriod.Status.REOPENED],
    ).update(needs_recalculation=True)


def parse_moment(day, value, tz_name, not_before=None):
    if value in (None, ""):
        return None
    tz = ZoneInfo(tz_name)
    if isinstance(value, datetime):
        moment = value if value.tzinfo else value.replace(tzinfo=tz)
    elif isinstance(value, time):
        moment = datetime.combine(day, value, tzinfo=tz)
    else:
        text = str(value).strip()
        if "T" in text or (len(text) > 5 and "-" in text):
            moment = datetime.fromisoformat(text.replace("Z", "+00:00"))
            if moment.tzinfo is None:
                moment = moment.replace(tzinfo=tz)
        else:
            parts = text.split(":")
            try:
                hour = int(parts[0])
                minute = int(parts[1]) if len(parts) > 1 else 0
            except (TypeError, ValueError) as exc:
                raise BusinessError("Use HH:MM for check-in and check-out times.") from exc
            moment = datetime.combine(day, time(hour, minute), tzinfo=tz)
    if not_before and moment <= not_before:
        moment += timedelta(days=1)
    return moment


def snapshot(record):
    return {
        "status": record.status,
        "check_in": record.check_in.isoformat() if record.check_in else None,
        "check_out": record.check_out.isoformat() if record.check_out else None,
        "late_minutes": record.late_minutes,
        "early_departure_minutes": record.early_departure_minutes,
        "working_minutes": record.working_minutes,
        "overtime_minutes": record.overtime_minutes,
        "break_minutes": record.break_minutes,
        "notes": record.notes,
        "shift_slot": record.shift_slot,
    }


def record_to_dict(record):
    return {
        "id": record.id,
        "status": record.status,
        "late_minutes": record.late_minutes,
        "early_departure_minutes": record.early_departure_minutes,
        "working_minutes": record.working_minutes,
        "overtime_minutes": record.overtime_minutes,
        "is_incomplete": record.is_incomplete,
        "check_in": record.check_in,
        "check_out": record.check_out,
        "notes": record.notes,
        "shift_slot": record.shift_slot,
        "break_minutes": record.break_minutes,
    }


def merge_slot_dicts(items):
    if not items:
        return None
    if len(items) == 1:
        return items[0]
    priority = [
        "incomplete",
        "absent",
        "half_day",
        "late",
        "present",
        "holiday_worked",
        "off_worked",
        "on_leave",
        "public_holiday",
        "weekly_off",
    ]
    status = items[0]["status"]
    for candidate in priority:
        if any(item["status"] == candidate for item in items):
            status = candidate
            break
    primary = min(items, key=lambda item: item["shift_slot"])
    merged = dict(primary)
    merged.update(
        {
            "status": status,
            "late_minutes": sum(item["late_minutes"] for item in items),
            "early_departure_minutes": max(item["early_departure_minutes"] for item in items),
            "working_minutes": sum(item["working_minutes"] for item in items),
            "overtime_minutes": sum(item["overtime_minutes"] for item in items),
            "is_incomplete": any(item["is_incomplete"] for item in items),
            "extra_slots": len(items) - 1,
        }
    )
    return merged


def load_attendance_map(employee_ids, start, end):
    grouped = defaultdict(lambda: defaultdict(list))
    rows = AttendanceRecord.objects.filter(employee_id__in=list(employee_ids), date__range=(start, end))
    for row in rows:
        grouped[row.employee_id][row.date].append(record_to_dict(row))
    merged = {}
    for employee_id, days in grouped.items():
        merged[employee_id] = {day: merge_slot_dicts(items) for day, items in days.items()}
    return merged


def local_hhmm(moment, tz_name):
    if not moment:
        return None
    try:
        tz = ZoneInfo(tz_name)
    except Exception:
        tz = ZoneInfo("UTC")
    return moment.astimezone(tz).strftime("%H:%M")


@transaction.atomic
def upsert_attendance(
    *,
    employee,
    day,
    actor,
    shift_slot=1,
    check_in=None,
    check_out=None,
    status=None,
    notes="",
    break_minutes=None,
    reason="",
    source=AttendanceRecord.Source.MANUAL,
    request=None,
    allow_without_reason=False,
    keep_existing_times=False,
):
    office = OfficeSettings.load()
    if not employed_on(employee, day):
        raise BusinessError(f"{employee.full_name} was not employed on {day.isoformat()}.")
    if employee.status not in WORKING_STATUSES and source != AttendanceRecord.Source.CORRECTION:
        if employee.exit_date and day > employee.exit_date:
            raise BusinessError("Attendance cannot be marked after the employee's exit date.")
    assert_month_editable(day)
    slot = int(shift_slot or 1)
    if slot < 1:
        raise BusinessError("Shift slot must be 1 or greater.")
    existing = AttendanceRecord.objects.filter(employee=employee, date=day, shift_slot=slot).select_for_update().first()
    if existing and not allow_without_reason and len((reason or "").strip()) < 3:
        raise BusinessError("Enter a reason of at least 3 characters before changing an existing attendance record.")
    if status and status not in VALID_STATUSES:
        raise BusinessError("That attendance status is not recognized.")

    template = shift_template(employee, office)
    offs = resolve_weekly_offs(employee, office)
    is_off = day.weekday() in offs
    is_holiday = day in holiday_dates_for(employee, day, day)
    leave_rows = leave_index([employee.id], day, day).get(employee.id, [])
    has_leave = covers_leave(leave_rows, day)
    parsed_in = parse_moment(day, check_in, template["tz"])
    parsed_out = parse_moment(day, check_out, template["tz"], not_before=parsed_in)
    if keep_existing_times and existing:
        parsed_in = parsed_in or existing.check_in
        parsed_out = parsed_out or existing.check_out
    break_value = template["break"] if break_minutes is None else int(break_minutes)
    if break_value < 0:
        raise BusinessError("Break minutes cannot be negative.")
    try:
        evaluated = evaluate_times(
            day=day,
            tz_name=template["tz"],
            shift_start=template["start"],
            shift_end=template["end"],
            overnight=template["overnight"],
            break_minutes=break_value,
            grace_minutes=template["grace"],
            early_grace_minutes=template["early"],
            half_day_threshold_minutes=office.half_day_threshold_minutes,
            check_in=parsed_in,
            check_out=parsed_out,
            is_weekly_off=is_off,
            is_holiday=is_holiday,
            has_leave=has_leave,
            holiday_work_is_overtime=office.holiday_work_counts_as_overtime,
            forced_status=status or None,
            today=office_now(office).date(),
        )
    except TimeError as exc:
        raise BusinessError(str(exc)) from exc
    final_status = evaluated["status"]
    if final_status == "unmarked":
        raise BusinessError("Add a check-in time or choose an attendance status.")

    previous = snapshot(existing) if existing else None
    record = existing or AttendanceRecord(employee=employee, date=day, shift_slot=slot, created_by=actor if getattr(actor, "is_authenticated", False) else None)
    record.shift = template["shift"]
    record.shift_start = template["start"]
    record.shift_end = template["end"]
    record.shift_overnight = evaluated["shift_overnight"]
    record.check_in = parsed_in
    record.check_out = parsed_out
    record.break_minutes = break_value
    record.scheduled_minutes = evaluated["scheduled_minutes"]
    record.working_minutes = evaluated["working_minutes"]
    record.late_minutes = evaluated["late_minutes"]
    record.early_departure_minutes = evaluated["early_departure_minutes"]
    record.overtime_minutes = evaluated["overtime_minutes"]
    record.status = final_status
    record.notes = notes if notes is not None else record.notes
    record.is_incomplete = evaluated["is_incomplete"]
    record.updated_by = actor if getattr(actor, "is_authenticated", False) else None
    if previous and (previous["status"] != record.status or previous["check_in"] != (record.check_in.isoformat() if record.check_in else None) or previous["check_out"] != (record.check_out.isoformat() if record.check_out else None)):
        record.source = AttendanceRecord.Source.CORRECTION
    else:
        record.source = source
    record.save()
    if previous:
        AttendanceAdjustment.objects.create(
            attendance=record,
            previous_data=previous,
            new_data=snapshot(record),
            reason=(reason or "").strip(),
            adjusted_by=actor if getattr(actor, "is_authenticated", False) else None,
        )
    audit(
        actor=actor,
        action="attendance.correct" if previous else "attendance.mark",
        instance=record,
        summary=f"{'Corrected' if previous else 'Marked'} {employee.full_name} on {day.isoformat()} as {record.get_status_display()}.",
        changes={"before": previous, "after": snapshot(record)} if previous else {"after": snapshot(record)},
        reason=reason or "",
        request=request,
    )
    return record


def bulk_mark(*, actor, day, employee_ids, check_in=None, check_out=None, status=None, notes="", reason="", request=None):
    employees = {item.id: item for item in visible_employees(actor).filter(id__in=employee_ids)}
    created = 0
    updated = 0
    errors = []
    for employee_id in employee_ids:
        employee = employees.get(int(employee_id))
        if employee is None:
            errors.append({"employee_id": employee_id, "detail": "Employee was not found or is outside your access."})
            continue
        if not can_manage_attendance(actor, employee):
            errors.append({"employee_id": employee_id, "detail": "You cannot mark attendance for this employee."})
            continue
        existed = AttendanceRecord.objects.filter(employee=employee, date=day, shift_slot=1).exists()
        try:
            upsert_attendance(
                employee=employee,
                day=day,
                actor=actor,
                check_in=check_in,
                check_out=check_out,
                status=status,
                notes=notes,
                reason=reason,
                request=request,
                allow_without_reason=not existed,
            )
        except BusinessError as exc:
            errors.append({"employee_id": employee_id, "detail": str(exc.detail)})
            continue
        if existed:
            updated += 1
        else:
            created += 1
    return {"created": created, "updated": updated, "errors": errors}


def employee_check_in(user, notes="", request=None):
    office = OfficeSettings.load()
    if not office.self_checkin_enabled:
        raise BusinessError("Employee check-in is turned off by the office settings.")
    employee = getattr(user, "employee_profile", None)
    if employee is None or employee.status not in WORKING_STATUSES:
        raise BusinessError("This account is not linked to an active employee profile.", status_code=403)
    now = office_now(office)
    existing = AttendanceRecord.objects.filter(employee=employee, date=now.date(), shift_slot=1).first()
    if existing and existing.check_in:
        raise BusinessError("You have already checked in today.")
    return upsert_attendance(
        employee=employee,
        day=now.date(),
        actor=user,
        check_in=now,
        check_out=existing.check_out if existing else None,
        notes=notes or (existing.notes if existing else ""),
        source=AttendanceRecord.Source.SELF,
        request=request,
        allow_without_reason=True,
        keep_existing_times=False,
    )


def employee_check_out(user, notes="", request=None):
    office = OfficeSettings.load()
    if not office.self_checkin_enabled:
        raise BusinessError("Employee check-out is turned off by the office settings.")
    employee = getattr(user, "employee_profile", None)
    if employee is None:
        raise BusinessError("This account is not linked to an employee profile.", status_code=403)
    now = office_now(office)
    existing = AttendanceRecord.objects.filter(employee=employee, date=now.date(), shift_slot=1).first()
    if existing is None or not existing.check_in:
        raise BusinessError("Check in before checking out.")
    if existing.check_out:
        raise BusinessError("You have already checked out today.")
    return upsert_attendance(
        employee=employee,
        day=now.date(),
        actor=user,
        check_in=existing.check_in,
        check_out=now,
        notes=notes or existing.notes,
        source=AttendanceRecord.Source.SELF,
        request=request,
        allow_without_reason=True,
    )


def materialize_employee(employee, year, month, today, actor=None):
    start, end = month_bounds(year, month)
    last = min(end, today - timedelta(days=1))
    if last < start:
        return 0
    office = OfficeSettings.load()
    offs = resolve_weekly_offs(employee, office)
    holidays = holiday_dates_for(employee, start, last)
    leave_rows = leave_index([employee.id], start, last).get(employee.id, [])
    existing_dates = set(
        AttendanceRecord.objects.filter(employee=employee, date__range=(start, last)).values_list("date", flat=True)
    )
    template = shift_template(employee, office)
    created = 0
    for offset in range((last - start).days + 1):
        day = start + timedelta(days=offset)
        if day in existing_dates or not employed_on(employee, day):
            continue
        try:
            evaluated = evaluate_times(
                day=day,
                tz_name=template["tz"],
                shift_start=template["start"],
                shift_end=template["end"],
                overnight=template["overnight"],
                break_minutes=template["break"],
                grace_minutes=template["grace"],
                early_grace_minutes=template["early"],
                half_day_threshold_minutes=office.half_day_threshold_minutes,
                check_in=None,
                check_out=None,
                is_weekly_off=day.weekday() in offs,
                is_holiday=day in holidays,
                has_leave=covers_leave(leave_rows, day),
                today=today,
            )
        except TimeError:
            continue
        if evaluated["status"] == "unmarked":
            continue
        AttendanceRecord.objects.create(
            employee=employee,
            date=day,
            shift_slot=1,
            shift=template["shift"],
            shift_start=template["start"],
            shift_end=template["end"],
            shift_overnight=evaluated["shift_overnight"],
            break_minutes=template["break"],
            scheduled_minutes=evaluated["scheduled_minutes"],
            status=evaluated["status"],
            source=AttendanceRecord.Source.SYSTEM,
            is_incomplete=False,
            created_by=actor if getattr(actor, "is_authenticated", False) else None,
        )
        created += 1
    return created


def apply_leave_attendance(leave_request):
    """Mark approved leave on working days without replacing a real punch."""
    office = OfficeSettings.load()
    employee = leave_request.employee
    offs = resolve_weekly_offs(employee, office)
    warnings = []
    day = leave_request.start_date
    while day <= leave_request.end_date:
        holiday = day in holiday_dates_for(employee, day, day)
        if day.weekday() in offs or holiday or not employed_on(employee, day):
            day += timedelta(days=1)
            continue
        existing = list(AttendanceRecord.objects.filter(employee=employee, date=day))
        punched = [row for row in existing if row.check_in or row.check_out]
        if punched:
            warnings.append(f"{day.isoformat()} already has a check-in, so leave did not replace it.")
        elif existing:
            for row in existing:
                if row.status in {"absent", "on_leave", "weekly_off", "public_holiday"} or row.source == AttendanceRecord.Source.SYSTEM:
                    row.status = AttendanceRecord.Status.ON_LEAVE
                    row.source = AttendanceRecord.Source.SYSTEM
                    row.save(update_fields=["status", "source", "updated_at"])
        else:
            template = shift_template(employee, office)
            AttendanceRecord.objects.create(
                employee=employee,
                date=day,
                shift=template["shift"],
                shift_start=template["start"],
                shift_end=template["end"],
                shift_overnight=template["overnight"],
                break_minutes=template["break"],
                status=AttendanceRecord.Status.ON_LEAVE,
                source=AttendanceRecord.Source.SYSTEM,
            )
        day += timedelta(days=1)
    return warnings


def clear_system_leave(employee, start, end):
    AttendanceRecord.objects.filter(
        employee=employee,
        date__range=(start, end),
        status=AttendanceRecord.Status.ON_LEAVE,
        source=AttendanceRecord.Source.SYSTEM,
        check_in__isnull=True,
        check_out__isnull=True,
    ).delete()


def import_workbook(upload, actor, request=None):
    try:
        workbook = load_workbook(upload, data_only=True)
    except Exception as exc:
        raise BusinessError("The file could not be read. Upload an .xlsx workbook.") from exc
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        raise BusinessError("The workbook is empty.")
    headers = [str(cell).strip().lower() if cell is not None else "" for cell in rows[0]]
    required = {"employee_code", "date"}
    if not required.issubset(set(headers)):
        raise BusinessError("The first row must include employee_code and date columns.")

    def col(name):
        return headers.index(name) if name in headers else None

    index = {name: col(name) for name in ("employee_code", "date", "check_in", "check_out", "status", "notes", "shift_slot", "break_minutes", "reason")}
    created = 0
    updated = 0
    errors = []
    employees = {item.employee_code.lower(): item for item in visible_employees(actor)}
    for row_number, row in enumerate(rows[1:], start=2):
        if not row or all(cell in (None, "") for cell in row):
            continue
        try:
            code = str(row[index["employee_code"]] or "").strip()
            employee = employees.get(code.lower())
            if employee is None:
                raise BusinessError(f"Unknown employee code {code}.")
            if not can_manage_attendance(actor, employee):
                raise BusinessError("You cannot import attendance for this employee.")
            raw_date = row[index["date"]]
            if hasattr(raw_date, "date"):
                day = raw_date.date() if isinstance(raw_date, datetime) else raw_date
            else:
                day = datetime.strptime(str(raw_date)[:10], "%Y-%m-%d").date()
            existed = AttendanceRecord.objects.filter(employee=employee, date=day, shift_slot=int(row[index["shift_slot"]] or 1) if index["shift_slot"] is not None else 1).exists()
            upsert_attendance(
                employee=employee,
                day=day,
                actor=actor,
                shift_slot=int(row[index["shift_slot"]] or 1) if index["shift_slot"] is not None and row[index["shift_slot"]] not in (None, "") else 1,
                check_in=row[index["check_in"]] if index["check_in"] is not None else None,
                check_out=row[index["check_out"]] if index["check_out"] is not None else None,
                status=(str(row[index["status"]]).strip() if index["status"] is not None and row[index["status"]] else None),
                notes=str(row[index["notes"]] or "") if index["notes"] is not None else "",
                break_minutes=int(row[index["break_minutes"]]) if index["break_minutes"] is not None and row[index["break_minutes"]] not in (None, "") else None,
                reason=str(row[index["reason"]] or "") if index["reason"] is not None else "",
                source=AttendanceRecord.Source.IMPORT,
                request=request,
                allow_without_reason=not existed,
            )
            if existed:
                updated += 1
            else:
                created += 1
        except BusinessError as exc:
            errors.append({"row": row_number, "detail": str(exc.detail)})
        except Exception as exc:
            errors.append({"row": row_number, "detail": str(exc)})
    return {"created": created, "updated": updated, "errors": errors}


def present_hours(record_dict):
    if not record_dict:
        return "0.00"
    return format(hours(record_dict.get("working_minutes") or 0), "f")
