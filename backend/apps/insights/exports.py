from io import BytesIO

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


HEADER_FILL = PatternFill("solid", fgColor="0F766E")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def xlsx_response(filename, sheets):
    workbook = Workbook()
    for index, sheet in enumerate(sheets):
        worksheet = workbook.active if index == 0 else workbook.create_sheet()
        worksheet.title = sheet["title"][:31]
        worksheet.append(sheet["headers"])
        for cell in worksheet[1]:
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
        for row in sheet["rows"]:
            worksheet.append(row)
        for column in worksheet.columns:
            worksheet.column_dimensions[column[0].column_letter].width = min(max(len(str(column[0].value or "")) + 2, 14), 36)
    buffer = BytesIO()
    workbook.save(buffer)
    response = HttpResponse(
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


def pdf_response(filename, title, headers, rows, landscape_page=False):
    buffer = BytesIO()
    document = SimpleDocTemplate(buffer, pagesize=landscape(A4) if landscape_page else A4, leftMargin=24, rightMargin=24, topMargin=28, bottomMargin=28)
    styles = getSampleStyleSheet()
    data = [headers] + [["" if value is None else str(value) for value in row] for row in rows]
    table = Table(data, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0f766e")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#d0d5dd")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f4f7f6")]),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("LEFTPADDING", (0, 0), (-1, -1), 4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    document.build([Paragraph(title, styles["Heading2"]), Spacer(1, 8), table])
    response = HttpResponse(buffer.getvalue(), content_type="application/pdf")
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


def choose_format(request, filename_base, title, headers, rows, landscape_page=False):
    export_format = request.query_params.get("format", "xlsx")
    if export_format == "pdf":
        return pdf_response(f"{filename_base}.pdf", title, headers, rows, landscape_page=landscape_page)
    return xlsx_response(f"{filename_base}.xlsx", [{"title": title[:31], "headers": headers, "rows": rows}])
